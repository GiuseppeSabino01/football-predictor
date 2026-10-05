from copy import deepcopy
from io import BytesIO

import pytest
import pandas as pd
from openpyxl import load_workbook

from fantasy.decision_center import best_rotation_pairs, recommend_lineup, simulate_purchase
from fantasy.catalog import normalize_catalog_dataframe
from fantasy.export import build_listone_excel, restore_listone_excel
from fantasy.mantra import MANTRA_FORMATIONS, best_lineup, eligible, formation_slots, mantra_roles, solve_lineup
from fantasy.official_catalog import catalog_fingerprint, parse_official_html
from fantasy.service import (
    add_purchase, auction_manager_summary, auction_managers, auction_price_board,
    auction_trade_analysis, create_league, new_workspace, normalize_workspace,
    record_auction_purchase, roster_summary, run_auction_multiverse,
    set_preferred_xi, set_scoring_system, simulate_auction_purchases,
    top_xi_for_formation, update_league_settings,
)


def player(index, mantra="Pc", classic="A", score=60):
    return {"id": f"p{index}", "player_id": f"p{index}", "name": f"Player {index}",
            "team": "INT", "role": classic, "mantra_role": mantra, "quote": 10,
            "fvm": 100, "mantra_fvm": 80, "price": 2, "fantasy_score": score,
            "expected_fantasy_average": 6.5, "starter_probability": 90,
            "reliability": 80, "expected_goals": 2, "expected_assists": 1}


def league(system="mantra", **kwargs):
    return create_league(new_workspace(), "Test", initial_budget=500, participants=2,
                         scoring_system=system, **kwargs)


def eleven():
    return [player(1, "Por", "P"), player(2, "Dd", "D"), player(3, "Dc", "D"),
            player(4, "Dc/B", "D"), player(5, "Ds/E", "D"), player(6, "M", "C"),
            player(7, "M/C", "C"), player(8, "C", "C"), player(9, "W/A", "A"),
            player(10, "W/A", "A"), player(11, "Pc", "A")]


def test_old_leagues_are_classic_after_normalization():
    old = league("classic")
    old.pop("scoring_system")
    assert normalize_workspace({"leagues": [old]})["leagues"][0]["scoring_system"] == "classic"


def test_switch_preserves_same_auction_and_every_purchase():
    draft = league("classic")
    record_auction_purchase(draft, auction_managers(draft)[0]["id"], player(1), 10)
    record_auction_purchase(draft, auction_managers(draft)[1]["id"], player(2), 12)
    saved = deepcopy(draft)
    set_scoring_system(draft, "mantra")
    assert draft["id"] == saved["id"]
    for key in ("purchases", "auction_managers", "auction_sale_events", "player_notes", "auction_player_tiers"):
        assert draft[key] == saved[key]
    assert draft["modifier_enabled"] is False
    set_scoring_system(draft, "classic")
    assert draft["purchases"] == saved["purchases"]
    assert draft["modifier_enabled"] == saved["modifier_enabled"]


def test_mantra_can_buy_more_than_classic_attack_quota():
    draft = league()
    for index in range(8):
        add_purchase(draft, player(index), 1)
    assert roster_summary(draft)["roster_size"] == 8
    before = deepcopy(draft)
    with pytest.raises(ValueError, match="Classic"):
        set_scoring_system(draft, "classic")
    assert draft == before


def test_mantra_limits_apply_to_opponents_and_remaining_budget():
    draft = league(mantra_roster_slots={"P": 1, "movement": 10})
    manager = auction_managers(draft)[1]["id"]
    for index in range(10):
        record_auction_purchase(draft, manager, player(index), 1)
    summary = auction_manager_summary(draft, manager)
    assert summary["remaining_slots"] == 1
    assert summary["target_size"] == 11
    assert summary["remaining_budget"] == 490
    with pytest.raises(ValueError, match="Mantra"):
        record_auction_purchase(draft, manager, player(11), 1)
    record_auction_purchase(draft, manager, player(12, "Por", "P"), 1)
    assert auction_manager_summary(draft, manager)["complete"]


def test_mantra_role_parser_and_unknown_roles():
    assert mantra_roles("dd;ds / e") == ("Dd", "Ds", "E")
    assert mantra_roles("p") == ("Por",)
    assert mantra_roles({"role": "D"}) == ()
    assert not eligible(player(1, ""), "Dc")


def test_all_eleven_official_modules_have_eleven_slots():
    assert len(MANTRA_FORMATIONS) == 11
    for module in MANTRA_FORMATIONS:
        assert len(formation_slots(module)) == 11
        assert formation_slots(module)[0] == "Por"


def test_solver_uses_multi_role_only_once_and_not_greedily():
    team = eleven()
    team[6] = player(7, "M/C", "C", score=500)
    selected = solve_lineup(team, "4-3-3")
    assert len(selected) == 11
    assert len({p["player_id"] for p in selected}) == 11
    assert all(eligible(p, p["lineup_slot"]) for p in selected)
    assert not best_lineup([player(i, "Pc") for i in range(11)])["complete"]


def test_mantra_requires_mediano_not_three_centrocampisti():
    team = eleven()
    for p in team:
        if p["role"] == "C":
            p["mantra_role"] = "C"
    assert len(solve_lineup(team, "4-3-3")) < 11


def test_top_eleven_and_validation_respect_mantra_positions():
    draft = league()
    draft["purchases"] = eleven()
    assert len(top_xi_for_formation(draft, "4-3-3")) == 11
    set_preferred_xi(draft, [p["player_id"] for p in eleven()], formation="4-3-3")
    draft["purchases"][5]["mantra_role"] = "Pc"
    draft["purchases"][6]["mantra_role"] = "Pc"
    with pytest.raises(ValueError, match="Mantra"):
        set_preferred_xi(draft, [p["player_id"] for p in eleven()], formation="4-3-3")


def test_top_eleven_auto_selects_a_valid_mantra_module():
    draft = league()
    draft["purchases"] = eleven()
    set_preferred_xi(draft, [p["player_id"] for p in eleven()])
    assert draft["preferred_formation"] in MANTRA_FORMATIONS
    assert len(top_xi_for_formation(draft, draft["preferred_formation"])) == 11


def test_imported_catalog_preserves_optional_mantra_columns():
    players = normalize_catalog_dataframe(pd.DataFrame([{
        "Giocatore": "Test", "Squadra": "INT", "Ruolo": "D", "RM": "Dd;Ds;E",
        "Quotazione": 10, "Quotazione Mantra": 11, "FVM Classic": 70, "FVM Mantra": 90,
    }]))
    assert players[0]["mantra_role"] == "Dd/Ds/E"
    assert players[0]["mantra_quote"] == 11
    assert players[0]["mantra_fvm"] == 90
    assert players[0]["fvm"] == 70


def test_matchday_assistant_does_not_return_a_classic_formation():
    draft = league()
    draft["purchases"] = eleven()
    result = recommend_lineup(draft, eleven(), matchday=1, news_items=[])
    assert result["complete"]
    assert all(eligible(p, p["lineup_slot"]) for p in result["players"])
    draft["purchases"] = [player(i) for i in range(11)]
    assert not recommend_lineup(draft, draft["purchases"], news_items=[])["complete"]


def test_incomplete_mantra_lineup_never_duplicates_selected_players_on_bench():
    draft = league()
    draft["purchases"] = eleven()[:5]
    result = recommend_lineup(draft, draft["purchases"], news_items=[])
    assert not result["complete"]
    assert result["formation"] in MANTRA_FORMATIONS
    assert not {p["player_id"] for p in result["players"]} & {p["player_id"] for p in result["bench"]}


def test_mantra_rotation_can_share_positions_across_classic_departments():
    draft = league()
    first, second = player(1, "W/T", "C"), player(2, "W/A", "A")
    first["team"], second["team"] = "INT", "MIL"
    draft["purchases"] = [first, second]
    pairs = best_rotation_pairs(draft, draft["purchases"], start_matchday=1)
    assert len(pairs) == 1
    assert pairs[0]["role"] == "W"


def test_mantra_price_uses_mantra_fvm_not_classic_fvm():
    draft = league()
    p = player(1)
    assert auction_price_board(draft, [p])[p["id"]]["initial"] == 40
    set_scoring_system(draft, "classic")
    assert auction_price_board(draft, [p])[p["id"]]["initial"] == 50


def test_html_preserves_both_official_prices_and_multiple_mantra_roles():
    html = '''<table><thead><tr><th>FVM</th></tr></thead><tbody>
    <tr data-filter-role-classic="d" data-filter-role-mantra="dd;ds;e">
    <th><a href="/serie-a/squadre/inter/test">Test</a></th><td>INT</td>
    <td data-col-key="c_qi">9</td><td data-col-key="c_qa">10</td><td data-col-key="c_fvm">70</td>
    <td data-col-key="m_qi">8</td><td data-col-key="m_qa">11</td><td data-col-key="m_fvm">90</td>
    </tr></tbody></table>'''
    parsed = parse_official_html(html, [player(1, classic="D")])[0]
    assert parsed["mantra_role"] == "Dd/Ds/E"
    assert parsed["fvm"] == 70
    assert parsed["mantra_fvm"] == 90
    assert parsed["quote"] == 10
    assert parsed["mantra_quote"] == 11


def test_fingerprint_detects_mantra_changes():
    p = player(1)
    assert catalog_fingerprint([p]) != catalog_fingerprint([{**p, "mantra_role": "A"}])


def test_excel_roundtrip_preserves_mantra_configuration_and_roles():
    draft = league()
    catalog = eleven()
    add_purchase(draft, catalog[0], 8)
    raw = build_listone_excel(catalog, draft)
    book = load_workbook(BytesIO(raw), read_only=True)
    assert "Ruolo Mantra" in next(book["Listone"].iter_rows(values_only=True))
    restore_listone_excel(raw, catalog, draft)
    assert draft["scoring_system"] == "mantra"
    assert draft["purchases"][0]["mantra_role"] == "Por"
    assert draft["purchases"][0]["price"] == 8


def test_simulation_and_what_if_use_mantra_total_capacity():
    draft = league()
    catalog = [player(i) for i in range(30)]
    generated = simulate_auction_purchases(draft, catalog, 7, seed=3)
    assert len(generated) == 14
    assert simulate_purchase(draft, player(100), 1)["valid"]


def test_mantra_trade_exclusions_and_no_classic_role_constraints():
    draft = league()
    draft["purchases"] = [player(i) for i in range(8)]
    rival = auction_managers(draft)[1]
    rival["purchases"] = [player(i + 20) for i in range(8)]
    draft["auction_trade_excluded_player_ids"] = ["p0"]
    result = auction_trade_analysis(draft, [], limit=10)
    assert result["ready"]
    assert result["trades"]
    for trade in result["trades"]:
        assert "p0" not in {p["player_id"] for p in trade["outgoing"]}
        assert 0 <= trade["user_improvement"] <= 15
        assert 0 <= trade["opponent_improvement"] <= 15
        assert trade["gain_gap"] <= 5
        assert len(trade["incoming"]) in (1, 2, 3)


def test_multiverse_never_runs_classic_rules_for_mantra():
    with pytest.raises(ValueError, match="Classic"):
        run_auction_multiverse(league(), [])


def test_mantra_settings_preserve_slot_counts_and_normalization():
    draft = league()
    update_league_settings(draft, name="Test", initial_budget=500, participants=2,
                          game_mode="auction", modifier_enabled=True, captain_enabled=False,
                          roster_slots={"P": 3, "D": 8, "C": 8, "A": 6},
                          mantra_roster_slots={"P": 2, "movement": 25})
    normalized = normalize_workspace({"leagues": [draft]})["leagues"][0]
    assert normalized["scoring_system"] == "mantra"
    assert normalized["mantra_roster_slots"] == {"P": 2, "movement": 25}
    assert roster_summary(normalized)["target_size"] == 27
    assert not normalized["modifier_enabled"]


def _mantra_ui_app(section="switch", empty=False):
    import streamlit as st
    from test_mantra import eleven
    from fantasy.service import create_league, new_workspace
    from fantasy.ui import (
        _render_scoring_system, _render_mantra_coverage, _render_top_xi_editor,
        _render_matchday_assistant, _render_create_form,
    )

    class MemoryStorage:
        remote_available = False

        def save(self, workspace):
            st.session_state["saved_system"] = workspace["leagues"][0]["scoring_system"]
            return False

    if "test_workspace" not in st.session_state:
        workspace = new_workspace()
        draft = create_league(workspace, "UI test", participants=2,
                             scoring_system="classic" if section == "switch" else "mantra")
        draft["id"] = "ui-test"
        draft["purchases"] = [] if empty else eleven()
        workspace["catalog"] = eleven()
        st.session_state["test_workspace"] = workspace
    workspace = st.session_state["test_workspace"]
    draft = workspace["leagues"][0]
    storage = MemoryStorage()
    if section == "switch":
        _render_scoring_system(workspace, draft, storage)
    elif section == "create":
        _render_create_form(workspace, storage, "ui-test")
    else:
        _render_mantra_coverage(draft, workspace["catalog"])
        _render_top_xi_editor(workspace, draft, storage)
        _render_matchday_assistant(workspace, draft, storage, workspace["catalog"], [])


def test_switch_widget_saves_same_auction_without_losing_roster():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_function(_mantra_ui_app).run()
    assert not app.exception
    app.radio[0].set_value("mantra").run()
    assert not app.exception
    assert app.session_state["saved_system"] == "mantra"
    draft = app.session_state["test_workspace"]["leagues"][0]
    assert draft["id"] == "ui-test"
    assert len(draft["purchases"]) == 11
    app.radio[0].set_value("classic").run()
    assert not app.exception
    assert app.session_state["saved_system"] == "classic"


@pytest.mark.parametrize("empty", [True, False])
def test_mantra_coverage_pitch_and_matchday_ui_render(empty):
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_function(_mantra_ui_app, args=("pitch", empty)).run()
    assert not app.exception
    assert app.selectbox[0].label == "Modulo Mantra"
    assert len(app.selectbox[0].options) == 11


def test_mantra_creation_form_uses_movement_capacity_not_classic_quotas():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_function(_mantra_ui_app, args=("create",)).run()
    app.radio[1].set_value("mantra").run()
    assert not app.exception
    labels = {widget.label for widget in app.number_input}
    assert "Portieri Mantra" in labels
    assert "Giocatori di movimento Mantra" in labels
    assert not {"D", "C", "A"} & labels
